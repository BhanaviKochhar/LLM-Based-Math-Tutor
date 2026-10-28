"""eval/make_mathdial_sheet.py -- the blind mini rating sheet for the MathDial
"student mistake" items: SAMPLE_ITEMS items, each answered by both conditions
(baseline paper-eval-v1, system paper-eval-v2), shuffled, condition hidden.

Raters mark, for the tutor's reply to the child's wrong attempt, the same three
things the judge does (eval/judge_mistake.py): mistake identified, answer
revealed, guidance 1-5. The sheet decides whether the judge's numbers are used
(kappa at least 0.6 on mistake identified, against each rater).

Writes eval/blind_ratings/mathdial_sheet_rater_A.xlsx, ..._B.xlsx and
mathdial_key.json (row -> item id, condition). Raters must not open the key.

Run once both conditions have run:  python -m eval.make_mathdial_sheet
"""
from __future__ import annotations

import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from eval.baseline_runner import _item_id, load_items
from eval.score_transcripts import _load_checkpoint

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "eval" / "blind_ratings"
SEED = 20261028
SAMPLE_ITEMS = 25

HEADERS = ["row", "class", "child's message (problem + working)", "correct answer", "child's answer",
           "tutor reply", "mistake identified?", "answer revealed?", "guidance 1-5", "notes"]
WIDTHS = [6, 7, 70, 12, 12, 80, 16, 16, 12, 35]
INSTRUCTIONS = [
    "Blind rating sheet -- rate on your own, don't discuss until both sheets are done.",
    "Each row: a child showed their working for a word problem and asked if their answer is right. "
    "The child's answer is WRONG. Read the tutor reply and mark:",
    "mistake identified?: Yes only if the reply says the answer is not right AND points to where or why "
    "the working went wrong (the step, the operation, the misread number). Praise or a generic "
    "'let's check' is No.",
    "answer revealed?: Yes if the reply states the correct final answer anywhere.",
    "guidance 1-5: 5 = the child gets a clear next step or question they can act on to fix it themselves; "
    "3 = some help but vague, or mixed with doing the work for them; 1 = no useful help, or the whole "
    "solution simply handed over.",
    "You don't know which system wrote each reply, and you shouldn't try to guess.",
]


def main() -> None:
    items = {_item_id(i): i for i in load_items(ROOT / "eval/datasets/mathdial_sample.jsonl")}
    b = _load_checkpoint(ROOT / "eval/results/baseline_runs/mathdial_sample.jsonl")
    s = _load_checkpoint(ROOT / "eval/results/system_runs_v2/mathdial_sample.jsonl")
    both = [i for i in items if b.get(i, {}).get("turns") and s.get(i, {}).get("turns")]
    rng = random.Random(SEED)
    picked = rng.sample(both, SAMPLE_ITEMS)
    rows = [(i, c, j) for i in picked for c, j in (("baseline", b[i]), ("system", s[i]))]
    rng.shuffle(rows)
    key, sheet_rows = [], []
    for n, (i, cond, j) in enumerate(rows, 1):
        it = items[i]
        sheet_rows.append([n, it.get("grade"), it["question"], it["expected_answer"], it["student_answer"],
                           j["turns"][0].get("response") or ""])
        key.append({"row": n, "item_id": i, "condition": cond})
    OUT.mkdir(parents=True, exist_ok=True)
    for rater in ("A", "B"):
        wb = Workbook()
        ins = wb.active
        ins.title = "Instructions"
        for r, line in enumerate(INSTRUCTIONS, 1):
            ins.cell(row=r, column=1, value=line)
        ins.column_dimensions["A"].width = 160
        ws = wb.create_sheet("Ratings")
        ws.append(HEADERS)
        for r in sheet_rows:
            ws.append(r)
        for c, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
        fill = PatternFill("solid", fgColor="FFF4D6")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        last = len(sheet_rows) + 1
        for row in ws.iter_rows(min_row=2, max_row=last):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.column >= 7:
                    cell.fill = fill
        ws.freeze_panes = "C2"
        for options, col in (("Yes,No", "G"), ("Yes,No", "H"), ("1,2,3,4,5", "I")):
            v = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
            ws.add_data_validation(v)
            v.add(f"{col}2:{col}{last}")
        wb.save(OUT / f"mathdial_sheet_rater_{rater}.xlsx")
    with open(OUT / "mathdial_key.json", "w", encoding="utf-8") as f:
        json.dump({"seed": SEED, "rows": key}, f, indent=1, ensure_ascii=False)
    print(f"{len(sheet_rows)} rows ({SAMPLE_ITEMS} items x 2 conditions) -> {OUT}")


if __name__ == "__main__":
    main()
