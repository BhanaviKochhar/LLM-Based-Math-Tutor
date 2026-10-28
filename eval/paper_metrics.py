"""eval/paper_metrics.py -- every number the paper reports, computed in one
place from the saved transcripts, judge verdicts and blind rating sheets
(no API calls). Writes one result file (eval/results/paper_metrics_*.json +
.md) that the tables script reads.

Conditions: baseline = eval/results/baseline_runs/ (tag paper-eval-v1),
system = eval/results/system_runs_v2/ (tag paper-eval-v2). Pairing is always
by item id; an item missing from either side is left out of paired tests.

Decisions this script encodes (agreed 2026-10-27):
  - Final-answer correctness comes from eval/score_transcripts.score_item.
    A reply the rules leave undecided is settled by the judge's committed
    answer (matched in code, never by the judge); if the judge also commits
    to nothing, or there is no verdict, it counts as not correct and the
    count is reported.
  - Rubric scores (1-5) are compared with the Wilcoxon signed-rank test,
    Holm-corrected across the four dimensions. Clarity and restraint are
    labelled exploratory (the judge missed the agreement gate on them).
  - SVAMP is reported on its own labels plus a row without the debatable
    labels (DEBATABLE_SVAMP).
  - The verifier's post-check on clock-time questions compared minutes with
    a clock time; for those items the check is recomputed against the
    item's clock-time ground truth.
  - Reading level and the pedagogy measures are computed from the reply
    text and reported comparatively (baseline vs system), not as absolute
    grade levels.

Run:  python -m eval.paper_metrics
"""
from __future__ import annotations

import json
import re
import statistics
from collections import Counter
from pathlib import Path

import numpy as np
from scipy.stats import wilcoxon

from eval import answer_extraction
from eval.baseline_runner import _item_id, load_items
from eval.component_eval.results import write_result, write_summary_markdown
from eval.rating_agreement import RUBRIC, cohen_kappa, load_sheet, norm, weighted_kappa
from eval.score_transcripts import (_load_checkpoint, _was_refused, expected_oos_behavior, ground_truth_candidates,
                                    mcnemar_exact_p, newcombe_paired_ci, score_item, wilson_ci)
from scripts.llm import scope, verifier

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "eval" / "datasets"
RUNS = ROOT / "eval" / "results"
JUDGE = RUNS / "judge_runs"
CURRICULUM = [f"curriculum_benchmark_class{g}" for g in range(1, 6)]
SVAMP_N = 500
# SVAMP items whose official answer contradicts the question, needs negative
# numbers or rests on an impossible premise (read item by item, 2026-10-26/27).
DEBATABLE_SVAMP = {f"svamp-{n:04d}" for n in (15, 28, 35, 37, 50, 154, 161, 204, 228, 290, 333, 342,
                                             351, 356, 417, 456, 498)}
EXPLORATORY_RUBRIC = {"clarity", "restraint"}


# --------------------------------------------------------------------------- loading
def items_of(stem: str) -> dict[str, dict]:
    path = DATA / (f"{stem}.json" if stem == "multi_turn" else f"{stem}.jsonl")
    return {_item_id(i): i for i in load_items(path)}


def runs(condition: str, stem: str) -> dict[str, dict]:
    folder = "baseline_runs" if condition == "baseline" else "system_runs_v2"
    return _load_checkpoint(RUNS / folder / f"{stem}.jsonl")


def verdicts(condition: str, stem: str) -> dict[str, dict]:
    path = JUDGE / f"{condition}_{stem}.jsonl"
    if not path.exists():
        return {}
    out = {}
    for line in open(path, encoding="utf-8"):
        if line.strip():
            r = json.loads(line)
            if r.get("ok"):
                out[r["id"]] = r["verdict"]
    return out


# --------------------------------------------------------------------------- statistics
def rate(k: int, n: int) -> dict:
    lo, hi = wilson_ci(k, n) if n else (None, None)
    return {"k": k, "n": n, "rate": k / n if n else None, "ci95": [lo, hi]}


def paired_binary(a: dict[str, bool], b: dict[str, bool]) -> dict:
    """a = baseline, b = system, keyed by item id; paired on shared ids."""
    ids = sorted(set(a) & set(b))
    both = sum(a[i] and b[i] for i in ids)
    only_a = sum(a[i] and not b[i] for i in ids)
    only_b = sum(b[i] and not a[i] for i in ids)
    neither = len(ids) - both - only_a - only_b
    diff, lo, hi = newcombe_paired_ci(both, only_a, only_b, neither)
    return {"n": len(ids), "baseline": rate(both + only_a, len(ids)), "system": rate(both + only_b, len(ids)),
            "both": both, "only_baseline": only_a, "only_system": only_b, "neither": neither,
            "diff_system_minus_baseline": diff, "diff_ci95": [lo, hi], "mcnemar_p": mcnemar_exact_p(only_a, only_b)}


def paired_scores(a: dict[str, float], b: dict[str, float]) -> dict:
    ids = sorted(set(a) & set(b))
    xa, xb = [a[i] for i in ids], [b[i] for i in ids]
    d = [y - x for x, y in zip(xa, xb)]
    p = None
    if ids and any(d):
        p = float(wilcoxon(xa, xb).pvalue)
    elif ids:
        p = 1.0
    return {"n": len(ids), "baseline_mean": float(np.mean(xa)) if ids else None,
            "system_mean": float(np.mean(xb)) if ids else None,
            "baseline_sd": float(np.std(xa, ddof=1)) if len(ids) > 1 else None,
            "system_sd": float(np.std(xb, ddof=1)) if len(ids) > 1 else None,
            "mean_diff": float(np.mean(d)) if ids else None, "wilcoxon_p": p}


def holm(rows: dict[str, dict], key: str = "wilcoxon_p") -> None:
    """Holm-adjust the p-values of a family of rows in place (adds p_holm)."""
    named = [{"label": k, "p": v.get(key)} for k, v in rows.items() if v.get(key) is not None]
    ordered = sorted(named, key=lambda r: r["p"])
    m, running = len(ordered), 0.0
    for i, r in enumerate(ordered):
        running = max(running, min(1.0, (m - i) * r["p"]))
        rows[r["label"]]["p_holm"] = running


def median_ci(values: list[float], seed: int = 20261027, n_boot: int = 2000) -> dict:
    if not values:
        return {"n": 0}
    arr = np.asarray(values, dtype=float)
    rng = np.random.default_rng(seed)
    meds = np.median(rng.choice(arr, size=(n_boot, len(arr)), replace=True), axis=1)
    return {"n": len(arr), "median": float(np.median(arr)), "median_ci95": [float(np.percentile(meds, 2.5)),
            float(np.percentile(meds, 97.5))], "p95": float(np.percentile(arr, 95)), "mean": float(arr.mean())}


# --------------------------------------------------------------------------- correctness
def final_correct(item: dict, journey: dict, verdict: dict | None) -> tuple[bool | None, str]:
    """(correct, how): how is rules | judge | undecided | not_scorable."""
    s = score_item(item, journey)
    if s.get("correct") is not None:
        return bool(s["correct"]), "rules"
    if not s.get("needs_llm_extraction") and not (s.get("note") or "").startswith("system hasn't revealed"):
        return None, "not_scorable"
    committed = str((verdict or {}).get("committed_answer") or "").strip()
    cands = ground_truth_candidates(item)
    if committed and committed.upper() != "NONE" and cands:
        return any(answer_extraction.match(committed, c, item.get("question") or "") for c in cands), "judge"
    return False, "undecided"


def correctness(stem: str, ids: set[str] | None = None) -> dict:
    items = items_of(stem)
    out, how = {}, {}
    for cond in ("baseline", "system"):
        rr, vv = runs(cond, stem), verdicts(cond, stem)
        out[cond], how[cond] = {}, Counter()
        for i, it in items.items():
            if (ids is not None and i not in ids) or i not in rr or not rr[i].get("turns"):
                continue
            c, h = final_correct(it, rr[i], vv.get(i))
            how[cond][h] += 1
            if c is not None:
                out[cond][i] = c
    return {"paired": paired_binary(out["baseline"], out["system"]),
            "decided_by": {k: dict(v) for k, v in how.items()}, "_per_item": out}


# --------------------------------------------------------------------------- out of scope and scope
def out_of_scope() -> dict:
    items = items_of("out_of_scope")
    res = {}
    for cond in ("baseline", "system"):
        rr, vv = runs(cond, "out_of_scope"), verdicts(cond, "out_of_scope")
        per = {}
        for i, it in items.items():
            if i not in rr or not rr[i].get("turns"):
                continue
            s = score_item(it, rr[i])
            c = s.get("correct")
            if c is None and s.get("needs_judge"):   # baseline reply to a gloss item
                v = vv.get(i)
                c = None if v is None else (v.get("teaches_advanced") is not True)
            if c is not None:
                per[i] = (bool(c), expected_oos_behavior(it), s.get("outcome"), bool(it.get("boundary_case")))
        res[cond] = per
    paired = paired_binary({i: v[0] for i, v in res["baseline"].items()}, {i: v[0] for i, v in res["system"].items()})
    by_expected = {}
    for exp in ("answer", "gloss", "decline"):
        by_expected[exp] = paired_binary({i: v[0] for i, v in res["baseline"].items() if v[1] == exp},
                                         {i: v[0] for i, v in res["system"].items() if v[1] == exp})
    boundary = paired_binary({i: v[0] for i, v in res["baseline"].items() if v[3]},
                             {i: v[0] for i, v in res["system"].items() if v[3]})
    outcomes = {c: {exp: dict(Counter(v[2] for v in res[c].values() if v[1] == exp))
                    for exp in ("answer", "gloss", "decline")} for c in res}
    return {"overall": paired, "by_expected_behaviour": by_expected, "boundary_cases": boundary,
            "outcome_counts": outcomes}


def refused(journey: dict) -> bool:
    return any(t.get("mode") in ("scope_refused", "safety_blocked") for t in journey.get("turns") or []) \
        or bool(_was_refused(journey))


def scope_rates() -> dict:
    """False refusals on in-scope questions; recall on beyond-Class-5 ones."""
    out = {}
    for cond in ("baseline", "system"):
        fr = {}
        for stem in CURRICULUM:
            rr = runs(cond, stem)
            fr[stem] = [refused(rr[i]) for i in items_of(stem) if i in rr and rr[i].get("turns")]
        oos_items, rr = items_of("out_of_scope"), runs(cond, "out_of_scope")
        probes = [refused(rr[i]) for i, it in oos_items.items()
                  if i in rr and rr[i].get("turns") and expected_oos_behavior(it) == "answer"]
        sv_items, rsv = items_of("svamp"), runs(cond, "svamp")
        sv = [refused(rsv[i]) for i in list(sv_items)[:SVAMP_N] if i in rsv and rsv[i].get("turns")]
        curr = [x for v in fr.values() for x in v]
        stopped = [refused(rr[i]) or (cond == "system" and score_item(it, rr[i]).get("outcome") == "gloss")
                   for i, it in oos_items.items() if i in rr and rr[i].get("turns")
                   and expected_oos_behavior(it) in ("decline", "gloss")]
        out[cond] = {"false_refusal_curriculum": rate(sum(curr), len(curr)),
                     "false_refusal_oos_probes": rate(sum(probes), len(probes)),
                     "false_refusal_svamp": rate(sum(sv), len(sv)),
                     "false_refusal_all_in_scope": rate(sum(curr) + sum(probes) + sum(sv),
                                                        len(curr) + len(probes) + len(sv)),
                     "recall_out_of_scope": rate(sum(stopped), len(stopped))}
    return out


def held_out() -> dict:
    """System v2 only: the held-out sets were written after the fixes."""
    res = {}
    op_items, rr = items_of("held_out_openers"), runs("system", "held_out_openers")
    got = Counter()
    for i, it in op_items.items():
        if i not in rr or not rr[i].get("turns"):
            continue
        mode = rr[i]["turns"][0].get("mode")
        observed = "CLARIFY" if mode == "clarify" else "ADVANCED" if refused(rr[i]) else "NEW"
        got[(it.get("expected"), observed)] += 1
    res["openers"] = {f"{e}->{o}": n for (e, o), n in sorted(got.items())}
    ae_items, rr = items_of("held_out_ambiguous_everyday"), runs("system", "held_out_ambiguous_everyday")
    done = [i for i in ae_items if i in rr and rr[i].get("turns")]
    reading = json.load(open(RUNS / "held_out_ambiguous_reading.json", encoding="utf-8"))["items"]
    res["ambiguous_everyday"] = {"n": len(done), "refused": rate(sum(refused(rr[i]) for i in done), len(done)),
                                 **{f"meets_expected_{k}": rate(sum(x["meets"] for x in reading if x["kind"] == k),
                                                                sum(x["kind"] == k for x in reading))
                                    for k in ("ambiguous", "everyday")},
                                 "note": "behaviour read item by item (eval/results/held_out_ambiguous_reading.json)"}
    ho_items, rr = items_of("held_out_out_of_scope"), runs("system", "held_out_out_of_scope")
    per = {i: score_item(it, rr[i]) for i, it in ho_items.items() if i in rr and rr[i].get("turns")}
    truly = [i for i in per if ho_items[i].get("truly_out_of_scope")]
    probes = [i for i in per if not ho_items[i].get("truly_out_of_scope")]
    res["out_of_scope"] = {"stopped_beyond_class_5": rate(sum(per[i]["outcome"] in ("refused", "gloss") for i in truly),
                                                          len(truly)),
                           "answered_in_scope_probes": rate(sum(per[i]["outcome"] != "refused" for i in probes),
                                                            len(probes))}
    return res


# --------------------------------------------------------------------------- judge and raters
def judge_rubric(stems: list[str]) -> dict:
    rows, key_idea, halluc = {}, {}, {}
    for dim in RUBRIC:
        a, b = {}, {}
        for stem in stems:
            va, vb = verdicts("baseline", stem), verdicts("system", stem)
            a.update({f"{stem}/{i}": v["rubric"][dim] for i, v in va.items()})
            b.update({f"{stem}/{i}": v["rubric"][dim] for i, v in vb.items()})
        rows[dim] = {**paired_scores(a, b), "exploratory": dim in EXPLORATORY_RUBRIC,
                     "score_counts": {"baseline": {str(s): sum(1 for x in a.values() if x == s) for s in range(1, 6)},
                                      "system": {str(s): sum(1 for x in b.values() if x == s) for s in range(1, 6)}}}
    holm(rows)
    for cond in ("baseline", "system"):
        ki = Counter()
        hf = []
        for stem in stems:
            for v in verdicts(cond, stem).values():
                if v.get("key_idea") in ("present", "partial", "absent"):
                    ki[v["key_idea"]] += 1
                hf.append(bool(v["hallucination"]["flag"]))
        key_idea[cond] = {"counts": dict(ki), "present": rate(ki["present"], sum(ki.values()))}
        halluc[cond] = rate(sum(hf), len(hf))
    return {"rubric": rows, "key_idea": key_idea,
            "judge_hallucination_flags_unvalidated": halluc}


def human_ratings() -> dict:
    """Two team members rated the same 150 replies blind. Reported jointly (the
    paper names no individual rater): hallucination flagged by both and by
    either; rubric scores averaged over the two; agreement between them."""
    key = {k["row"]: k for k in json.load(open(ROOT / "eval/blind_ratings/rating_key.json", encoding="utf-8"))["rows"]}
    sheets = [{int(x["row"]): x for x in load_sheet(r)} for r in ("A", "B")]
    by = {"baseline": {}, "system": {}}
    for row, k in key.items():
        by[k["condition"]][k["item_id"]] = [s[row] for s in sheets]
    flags = {c: {i: [norm(r["halluc"]) == "yes" for r in rs] for i, rs in by[c].items()} for c in by}
    out = {"hallucination_both": paired_binary({i: all(v) for i, v in flags["baseline"].items()},
                                               {i: all(v) for i, v in flags["system"].items()}),
           "hallucination_either": paired_binary({i: any(v) for i, v in flags["baseline"].items()},
                                                 {i: any(v) for i, v in flags["system"].items()})}
    rub = {}
    for dim in RUBRIC:
        mean = {c: {i: sum(float(r[dim]) for r in rs) / 2 for i, rs in by[c].items()
                    if all(str(r[dim]).strip() for r in rs)} for c in by}
        rub[dim] = paired_scores(mean["baseline"], mean["system"])
    holm(rub)
    out["rubric_mean_of_two"] = rub
    a, b = ([s[row] for row in sorted(key)] for s in sheets)
    out["agreement"] = {"hallucination_kappa": cohen_kappa([norm(x["halluc"]) for x in a],
                                                           [norm(x["halluc"]) for x in b]),
                        **{f"{d}_weighted_kappa": weighted_kappa([x[d] for x in a], [x[d] for x in b])
                           for d in RUBRIC}}
    out["note"] = ("75 items not used in any earlier rating round; baseline replies from paper-eval-v1, "
                   "system replies from paper-eval-v2")
    return out


# --------------------------------------------------------------------------- MathDial (reply to a wrong attempt)
_MD_COLS = ["row", "class", "msg", "correct_answer", "child_answer", "reply", "mistake", "revealed", "guidance",
            "notes"]


def _mathdial_sheet(rater: str) -> dict[int, dict]:
    from openpyxl import load_workbook
    ws = load_workbook(ROOT / f"eval/blind_ratings/mathdial_sheet_rater_{rater}.xlsx", data_only=True)["Ratings"]
    return {int(x[0]): dict(zip(_MD_COLS, x)) for x in ws.iter_rows(min_row=2, values_only=True)
            if x and x[0] is not None}


def mathdial() -> dict:
    """The tutor's first reply to a child's wrong working. Answer revealed is
    checked in code (does the reply commit to the correct answer); the judge's
    mistake-identified, answer-revealed and guidance are used only where it
    agrees with the two team members (kappa >= 0.6 against each)."""
    items = items_of("mathdial_sample")
    jv = {c: verdicts(c, "mathdial_sample") for c in ("baseline", "system")}
    res = {}
    # Extractor check: the reply commits to the correct final answer.
    rev = {}
    for c in ("baseline", "system"):
        rr = runs(c, "mathdial_sample")
        rev[c] = {i: answer_extraction.judge(answer_extraction.committed_answer(rr[i]["turns"][0].get("response") or ""),
                                             ground_truth_candidates(it), it.get("question") or "") is True
                  for i, it in items.items() if i in rr and rr[i].get("turns")}
    res["answer_revealed_extractor"] = paired_binary(rev["baseline"], rev["system"])
    res["system_first_turn_modes"] = dict(Counter(runs("system", "mathdial_sample")[i]["turns"][0].get("mode")
                                                  for i in items if i in runs("system", "mathdial_sample")))
    # Humans on the 25-item sheet, and judge agreement with each.
    key = {k["row"]: k for k in json.load(open(ROOT / "eval/blind_ratings/mathdial_key.json", encoding="utf-8"))["rows"]}
    sheets = [_mathdial_sheet(r) for r in ("A", "B")]
    yn = lambda v: str(v or "").strip().lower() == "yes"
    agree = {}
    for field, jf in (("mistake", "mistake_identified"), ("revealed", "answer_revealed")):
        hum = [[yn(s[row][field]) for row in sorted(key)] for s in sheets]
        jud = [bool(jv[key[row]["condition"]].get(key[row]["item_id"], {}).get(jf)) for row in sorted(key)]
        agree[jf] = {"humans_kappa": cohen_kappa(*hum),
                     "judge_vs_member_kappas": [cohen_kappa(jud, h) for h in hum]}
    hg = [[s[row]["guidance"] for row in sorted(key)] for s in sheets]
    jg = [jv[key[row]["condition"]].get(key[row]["item_id"], {}).get("guidance") for row in sorted(key)]
    agree["guidance"] = {"humans_weighted_kappa": weighted_kappa(*hg),
                         "judge_vs_member_weighted_kappas": [weighted_kappa(jg, h) for h in hg]}
    res["agreement"] = agree
    trusted = {f: all(k is not None and k >= 0.6 for k in agree[f]["judge_vs_member_kappas"])
               for f in ("mistake_identified", "answer_revealed")}
    res["judge_trusted"] = trusted
    # Full-scale judge numbers (reported as main results only where trusted).
    for f in ("mistake_identified", "answer_revealed"):
        res[f"judge_{f}"] = paired_binary({i: bool(v[f]) for i, v in jv["baseline"].items()},
                                          {i: bool(v[f]) for i, v in jv["system"].items()})
    res["judge_guidance_exploratory"] = paired_scores({i: v["guidance"] for i, v in jv["baseline"].items()},
                                                      {i: v["guidance"] for i, v in jv["system"].items()})
    # Human numbers on the 25 rated items: identified by both / either member.
    hb = {"baseline": {}, "system": {}}
    for row, k in key.items():
        hb[k["condition"]][k["item_id"]] = [yn(s[row]["mistake"]) for s in sheets]
    res["humans_mistake_identified_both"] = paired_binary({i: all(v) for i, v in hb["baseline"].items()},
                                                          {i: all(v) for i, v in hb["system"].items()})
    res["humans_mistake_identified_either"] = paired_binary({i: any(v) for i, v in hb["baseline"].items()},
                                                            {i: any(v) for i, v in hb["system"].items()})
    return res


# --------------------------------------------------------------------------- multi-turn
def _expected_mode(text: str) -> str | None:
    """The dataset's expected action, as one controller mode. Labels that start a
    new problem (new_question, math_followup, new episode) are a fresh
    teach_invite; a give-up "directly explain" is a reveal; a conversation
    ending is not an action (None, left out)."""
    t = (text or "").lower()
    if "termination" in t:
        return None
    if t.startswith(("new_question", "math_followup", "new episode")):
        return "teach_invite"
    if "give_up" in t:
        return "reveal"
    for key, mode in (("reveal", "reveal"), ("diagnose_correct", "diagnose_correct"),
                      ("diagnose_wrong", "diagnose_wrong"), ("hint", "hint"), ("ack_conceptual", "ack_conceptual"),
                      ("teach_invite", "teach_invite"), ("gloss", "gloss"), ("co_solve", "co_solve"),
                      ("redirect", "redirect"), ("clarify", "clarify"), ("confusion", "confusion_help"),
                      ("correction", "correction_repair"), ("refus", "scope_refused")):
        if key in t:
            return mode
    return t.split(",")[0].strip()


def multi_turn() -> dict:
    """Controller accuracy: does each system turn take the action the dataset
    expects (system only -- the baseline has no actions)? Plus the best-effort
    final-answer score for both conditions."""
    convs = {c["conversation_id"]: c for c in json.load(open(DATA / "multi_turn.json", encoding="utf-8"))["conversations"]}
    actions = _action_match(convs, runs("system", "multi_turn"))
    final = {}
    for cond in ("baseline", "system"):
        r = runs(cond, "multi_turn")
        final[cond] = {cid: bool(score_item(c, r[cid])["correct"]) for cid, c in convs.items()
                       if cid in r and r[cid].get("turns") and score_item(c, r[cid]).get("correct") is not None}
    return {"controller_action_match": actions["overall"], "by_expected_action": actions["by_expected"],
            "final_answer_best_effort": paired_binary(final["baseline"], final["system"])}


def _action_match(convs: dict, rr: dict) -> dict:
    match = Counter()
    by_expected: dict[str, Counter] = {}
    for cid, c in convs.items():
        j = rr.get(cid)
        if not j or not j.get("turns"):
            continue
        for exp_t, got in zip(c["turns"], j["turns"]):
            exp = _expected_mode(exp_t.get("expected_controller_action") or "")
            if exp is None:
                continue
            mode = got.get("mode") or "gloss"
            ok = mode == exp or (exp == "reveal" and mode == "co_solve") or (exp == "co_solve" and mode == "reveal")
            match[ok] += 1
            by_expected.setdefault(exp, Counter())[ok] += 1
    return {"overall": rate(match[True], match[True] + match[False]),
            "by_expected": {k: rate(v[True], v[True] + v[False]) for k, v in sorted(by_expected.items())}}


# --------------------------------------------------------------------------- production latency
def production_latency() -> dict:
    """The live app's own calls in the pilot snapshot (2026-07-20 to 2026-08-22),
    from data/llm_runs.jsonl: tutoring calls only, before any evaluation run."""
    rows = []
    for line in open(ROOT / "data" / "llm_runs.jsonl", encoding="utf-8"):
        if not line.strip():
            continue
        r = json.loads(line)
        if r.get("model") in ("tutor-pipeline", "tutor-turn") and "2026-07-20" <= r["ts"][:10] <= "2026-08-22":
            rows.append(r)
    ok = [r for r in rows if r["result"].get("ok") and r["result"].get("latency_s") is not None]
    providers = Counter(r["result"].get("provider") for r in ok)
    slowest = sorted(ok, key=lambda r: -r["result"]["latency_s"])[:3]
    return {"calls": len(rows), "succeeded": len(ok), "latency_s": median_ci([r["result"]["latency_s"] for r in ok]),
            "providers": dict(providers),
            "fallback_share": rate(sum(v for k, v in providers.items() if k != "groq-direct"), len(ok)),
            "slowest": [{"ts": r["ts"], "latency_s": r["result"]["latency_s"],
                         "provider": r["result"].get("provider")} for r in slowest]}


# --------------------------------------------------------------------------- reading level and pedagogy
_MD_RE = re.compile(r"[*_#`>|]|\\[()\[\]]|^\s*[-•]\s*", re.MULTILINE)
_WORD_RE = re.compile(r"[A-Za-z]+(?:'[A-Za-z]+)?")
_SENT_SPLIT_RE = re.compile(r"(?<=[.!?])\s+|\n+")
_ENCOURAGE_RE = re.compile(r"\b(great job|well done|good job|nice work|you did it|awesome|excellent|"
                           r"keep it up|you can do it|great work|fantastic|brilliant)\b", re.IGNORECASE)
_INVITE_RE = re.compile(r"\b(your turn|give it a (?:try|go)|try (?:it|this|again)|can you|what do you think|"
                        r"have a go|now you try)\b", re.IGNORECASE)
_STEP_RE = re.compile(r"^\s*(?:\*\*)?(?:step\s*\d+|\d+[.)])", re.IGNORECASE | re.MULTILINE)
_EXAMPLE_RE = re.compile(r"\b(for example|for instance|imagine|e\.g\.|everyday example|think of)\b", re.IGNORECASE)
_PRACTICE_RE = re.compile(r"\b(practice|your turn|try this one)\b", re.IGNORECASE)
_BEYOND_RES = [re.compile(p, re.IGNORECASE) for p, _ in scope.BEYOND_PRIMARY_CONCEPTS]


def _syllables(word: str) -> int:
    w = word.lower()
    groups = re.findall(r"[aeiouy]+", w)
    n = len(groups) - (1 if w.endswith("e") and len(groups) > 1 and not w.endswith("le") else 0)
    return max(1, n)


def reading_level(text: str) -> dict | None:
    """Flesch-Kincaid grade, Coleman-Liau index, words per sentence and the
    share of words with 3+ syllables, on prose only (markdown and numbers
    stripped). Comparative use only: the formulas were built for US prose."""
    clean = _MD_RE.sub(" ", text or "")
    sentences = [s for s in _SENT_SPLIT_RE.split(clean) if _WORD_RE.search(s)]
    words = _WORD_RE.findall(clean)
    if len(words) < 10 or not sentences:
        return None
    syll = [_syllables(w) for w in words]
    nw, ns = len(words), len(sentences)
    letters = sum(len(w) for w in words)
    return {"fk_grade": 0.39 * nw / ns + 11.8 * sum(syll) / nw - 15.59,
            "coleman_liau": 0.0588 * (letters / nw * 100) - 0.296 * (ns / nw * 100) - 15.8,
            "words_per_sentence": nw / ns, "polysyllabic_share": sum(s >= 3 for s in syll) / nw, "words": nw}


def pedagogy(text: str) -> dict:
    t = text or ""
    sentences = [s for s in _SENT_SPLIT_RE.split(t) if s.strip()]
    return {"questions_to_child": sum(s.strip().endswith("?") for s in sentences),
            "invites_attempt": bool(_INVITE_RE.search(t)), "numbered_steps": len(_STEP_RE.findall(t)),
            "has_example": bool(_EXAMPLE_RE.search(t)), "encouragement": len(_ENCOURAGE_RE.findall(t)),
            "practice_question": bool(_PRACTICE_RE.search(t)),
            "beyond_class_5_terms": sum(bool(r.search(t)) for r in _BEYOND_RES)}


def text_measures(stems: list[str], limit_ids: dict[str, set] | None = None) -> dict:
    """First reply (teaching stance) and final reply, per condition, plus the
    early-reveal rate: the first reply already commits to the right answer."""
    out = {}
    for cond in ("baseline", "system"):
        agg = {"first": [], "final": []}
        ped = {"first": [], "final": []}
        reveal = []
        for stem in stems:
            items, rr = items_of(stem), runs(cond, stem)
            for i, it in items.items():
                if limit_ids and stem in limit_ids and i not in limit_ids[stem]:
                    continue
                j = rr.get(i)
                if not j or not j.get("turns") or refused(j):
                    continue
                first, final = j["turns"][0].get("response") or "", j["turns"][-1].get("response") or ""
                for k, txt in (("first", first), ("final", final)):
                    rl = reading_level(txt)
                    if rl:
                        agg[k].append(rl)
                    ped[k].append(pedagogy(txt))
                cands = ground_truth_candidates(it)
                if cands:
                    v = answer_extraction.judge(answer_extraction.committed_answer(first), cands, it.get("question") or "")
                    reveal.append(v is True)
        res = {"early_reveal_in_first_reply": rate(sum(reveal), len(reveal))}
        for k in ("first", "final"):
            res[f"reading_{k}"] = {m: median_ci([r[m] for r in agg[k]])
                                   for m in ("fk_grade", "coleman_liau", "words_per_sentence", "polysyllabic_share",
                                             "words")}
            p = ped[k]
            res[f"pedagogy_{k}"] = {
                "n": len(p),
                "mean_questions_to_child": statistics.mean(x["questions_to_child"] for x in p) if p else None,
                "invites_attempt": rate(sum(x["invites_attempt"] for x in p), len(p)),
                "has_numbered_steps": rate(sum(x["numbered_steps"] > 0 for x in p), len(p)),
                "has_example": rate(sum(x["has_example"] for x in p), len(p)),
                "has_encouragement": rate(sum(x["encouragement"] > 0 for x in p), len(p)),
                "practice_question": rate(sum(x["practice_question"] for x in p), len(p)),
                "mentions_beyond_class_5": rate(sum(x["beyond_class_5_terms"] > 0 for x in p), len(p))}
        out[cond] = res
    return out


# --------------------------------------------------------------------------- verifier and latency
def verifier_stats(stems: list[str]) -> dict:
    n_items = injected = 0
    wrong_values: list[str] = []
    post = Counter()
    # Coverage on items that have a checkable expression, per set, and how often the
    # verifier engaged on items that have none (conceptual, comparison, ...).
    coverage: dict[str, Counter] = {}
    for stem in stems:
        items, rr = items_of(stem), runs("system", stem)
        cov = coverage.setdefault(stem.replace("curriculum_benchmark_", ""), Counter())
        for i, it in items.items():
            j = rr.get(i)
            if not j or not j.get("turns"):
                continue
            n_items += 1
            checks = [t["verify"] for t in j["turns"] if t.get("verify") and t["verify"].get("verifiable")]
            checkable = bool(it.get("trusted_expression")) and it.get("verifiable", True) is not False
            cov[("checkable" if checkable else "not_checkable") + ("_engaged" if checks else "_missed")] += 1
            if not checks:
                continue
            injected += 1
            v = checks[-1]
            match = v.get("match")
            # Was the trusted value itself right? (Clock items log minutes here, not
            # the injected clock time, so they are left out of this count.)
            cands = ground_truth_candidates(it)
            if cands and v.get("computed") and not verifier.asks_for_clock_time(it.get("question") or ""):
                right = any(answer_extraction.match(str(v["computed"]), c, it.get("question") or "") for c in cands)
                post["value_right" if right else "value_wrong"] += 1
                if not right:
                    wrong_values.append(f"{stem}/{i}")
            if match is False and verifier.asks_for_clock_time(it.get("question") or ""):
                # The post-check compared minutes with a clock time; recheck the reply
                # against the item's clock-time ground truth instead.
                c = answer_extraction.judge(answer_extraction.committed_answer(j["turns"][-1].get("response") or ""),
                                            ground_truth_candidates(it), it.get("question") or "")
                match = c
                post["clock_rechecked"] += 1
            post[{True: "match", False: "mismatch", None: "no_number"}[match]] += 1
    def share(sets: list[str], kind: str) -> dict:
        c = sum((coverage[s] for s in sets), Counter())
        return rate(c[f"{kind}_engaged"], c[f"{kind}_engaged"] + c[f"{kind}_missed"])
    curriculum_sets = [s for s in coverage if s != "svamp"]
    return {"items": n_items, "with_trusted_value": rate(injected, n_items),
            "coverage_checkable": {**{s: share([s], "checkable") for s in coverage},
                                   "curriculum": share(curriculum_sets, "checkable")},
            "engaged_not_checkable_curriculum": share(curriculum_sets, "not_checkable"),
            "post_check": dict(post),
            "trusted_value_right": rate(post["value_right"], post["value_right"] + post["value_wrong"]),
            "items_with_wrong_trusted_value": wrong_values}


def latency(stem: str, ids: list[str] | None = None) -> dict:
    out = {}
    for cond in ("baseline", "system"):
        rr = runs(cond, stem)
        turns = [t for i, j in rr.items() if (ids is None or i in ids) and j.get("turns") for t in j["turns"]
                 if t.get("turn_wall_s") is not None]
        clean = [t for t in turns if not t.get("noisy") and not t.get("rate_limited_attempts")]
        out[cond] = {"turns": len(turns), "clean_turns": len(clean),
                     "clean_turn_wall_s": median_ci([t["turn_wall_s"] for t in clean]),
                     "clean_ttft_s": median_ci([t["ttft_s"] for t in clean if t.get("ttft_s") is not None]),
                     "all_turn_wall_s": median_ci([t["turn_wall_s"] for t in turns])}
    return out



# --------------------------------------------------------------------------- main
def _pct(r: dict) -> str:
    return "n/a" if r.get("rate") is None else f"{100 * r['rate']:.1f}% ({r['k']}/{r['n']})"


def main() -> None:
    svamp_ids = set(list(items_of("svamp"))[:SVAMP_N])
    m: dict = {}
    m["curriculum"] = {stem: correctness(stem) for stem in CURRICULUM}
    allb = {f"{s}/{i}": v for s in CURRICULUM for i, v in m["curriculum"][s]["_per_item"]["baseline"].items()}
    alls = {f"{s}/{i}": v for s in CURRICULUM for i, v in m["curriculum"][s]["_per_item"]["system"].items()}
    m["curriculum_all"] = paired_binary(allb, alls)
    sv = correctness("svamp", svamp_ids)
    m["svamp"] = {"official_labels": sv["paired"], "decided_by": sv["decided_by"],
                  "without_debatable": paired_binary(
                      {i: v for i, v in sv["_per_item"]["baseline"].items() if i not in DEBATABLE_SVAMP},
                      {i: v for i, v in sv["_per_item"]["system"].items() if i not in DEBATABLE_SVAMP})}
    m["out_of_scope"] = out_of_scope()
    m["scope"] = scope_rates()
    m["held_out"] = held_out()
    m["judge_curriculum"] = judge_rubric(CURRICULUM)
    m["judge_out_of_scope"] = judge_rubric(["out_of_scope"])
    if verdicts("system", "multi_turn"):
        m["judge_multi_turn"] = judge_rubric(["multi_turn"])
    m["human_ratings"] = human_ratings()
    m["mathdial"] = mathdial()
    m["multi_turn"] = multi_turn()
    m["text_curriculum"] = text_measures(CURRICULUM)
    m["text_svamp"] = text_measures(["svamp"], {"svamp": svamp_ids})
    m["verifier"] = verifier_stats(CURRICULUM + ["svamp"])
    m["latency_svamp"] = latency("svamp", sorted(svamp_ids))
    m["latency_system_v2_all_sets"] = {stem: latency(stem)["system"] for stem in CURRICULUM + ["out_of_scope"]}
    m["production_latency"] = production_latency()
    for stem in CURRICULUM:
        m["curriculum"][stem].pop("_per_item")

    lines = ["# Paper metrics", "", "Baseline = paper-eval-v1 transcripts; system = paper-eval-v2.", ""]
    c = m["curriculum_all"]
    lines += [f"- Curriculum final-answer accuracy: baseline {_pct(c['baseline'])}, system {_pct(c['system'])}, "
              f"diff {c['diff_system_minus_baseline']:+.3f} [{c['diff_ci95'][0]:+.3f}, {c['diff_ci95'][1]:+.3f}], "
              f"McNemar p = {c['mcnemar_p']:.3f}"]
    for label, r in (("SVAMP (official labels)", m["svamp"]["official_labels"]),
                     ("SVAMP (without debatable)", m["svamp"]["without_debatable"]),
                     ("Out-of-scope correct behaviour", m["out_of_scope"]["overall"])):
        lines.append(f"- {label}: baseline {_pct(r['baseline'])}, system {_pct(r['system'])}, "
                     f"McNemar p = {r['mcnemar_p']:.3f}")
    s = m["scope"]["system"]
    lines += [f"- System false refusals (all in-scope): {_pct(s['false_refusal_all_in_scope'])}; "
              f"recall on out-of-scope: {_pct(s['recall_out_of_scope'])}", "", "## Judge rubric (curriculum)"]
    for dim, r in m["judge_curriculum"]["rubric"].items():
        tag = " (exploratory)" if r["exploratory"] else ""
        lines.append(f"- {dim}{tag}: baseline {r['baseline_mean']:.2f}, system {r['system_mean']:.2f}, "
                     f"Wilcoxon p = {r['wilcoxon_p']:.3g}, Holm {r.get('p_holm', float('nan')):.3g} (n = {r['n']})")
    hr = m["human_ratings"]
    lines += ["", "## Two team members (150 blind replies)"]
    for label, k in (("hallucination flagged by both", "hallucination_both"),
                     ("hallucination flagged by either", "hallucination_either")):
        h = hr[k]
        lines.append(f"- {label}: baseline {_pct(h['baseline'])}, system {_pct(h['system'])}, "
                     f"McNemar p = {h['mcnemar_p']:.3f}")
    lines.append(f"- agreement: {', '.join(f'{k} {v:.2f}' for k, v in hr['agreement'].items() if v is not None)}")
    md = m["mathdial"]
    lines += ["", "## MathDial (reply to a wrong attempt)",
              f"- answer revealed (code): baseline {_pct(md['answer_revealed_extractor']['baseline'])}, "
              f"system {_pct(md['answer_revealed_extractor']['system'])}",
              f"- mistake identified (judge{'' if md['judge_trusted']['mistake_identified'] else ', not trusted'}): "
              f"baseline {_pct(md['judge_mistake_identified']['baseline'])}, "
              f"system {_pct(md['judge_mistake_identified']['system'])}",
              f"- judge agreement: {md['agreement']}",
              f"- system first-turn modes: {md['system_first_turn_modes']}"]
    mt = m["multi_turn"]
    lines += ["", f"## Multi-turn: controller action match {_pct(mt['controller_action_match'])}"]
    ver = m["verifier"]
    lines += ["", "## Verifier (system)",
              "- checkable items given a trusted value: " + ", ".join(f"{s} {_pct(r)}" for s, r in ver["coverage_checkable"].items()),
              f"- engaged on curriculum items with no checkable answer: {_pct(ver['engaged_not_checkable_curriculum'])}",
              f"- trusted value correct: {_pct(ver['trusted_value_right'])}; wrong on {', '.join(ver['items_with_wrong_trusted_value'])}"]
    pl = m["production_latency"]
    lines += ["", f"## Production latency (pilot): median {pl['latency_s'].get('median', float('nan')):.2f} s, "
                  f"P95 {pl['latency_s'].get('p95', float('nan')):.2f} s, {pl['succeeded']}/{pl['calls']} calls"]
    t = m["text_curriculum"]
    lines += ["", "## Reading level and pedagogy (curriculum, first reply)"]
    for cond in ("baseline", "system"):
        r = t[cond]
        lines.append(f"- {cond}: FK grade {r['reading_first']['fk_grade'].get('median', float('nan')):.1f}, "
                     f"words {r['reading_first']['words'].get('median', float('nan')):.0f}, "
                     f"early reveal {_pct(r['early_reveal_in_first_reply'])}, "
                     f"invites an attempt {_pct(r['pedagogy_first']['invites_attempt'])}")
    print("\n".join(lines))
    path = write_result("paper_metrics", dataset_sources=[{"path": str(DATA)}],
                        config={"svamp_n": SVAMP_N, "debatable_svamp": sorted(DEBATABLE_SVAMP)},
                        per_example_results=[], aggregate_metrics=m,
                        notes="All paper numbers; see the module docstring for the decisions encoded.")
    write_summary_markdown(path, lines)
    print(f"\nRaw: {path}")


if __name__ == "__main__":
    main()
